"""Official Shiller workbook adapter; publication is explicitly permission-gated."""
import copy
import io
import math
import time
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit

import httpx

PAGE_URL = "https://shillerdata.com/"
MAX_BYTES = 10 * 1024 * 1024
ALIASES = {"P": "SH_P", "D": "SH_D", "E": "SH_E", "CPI": "SH_CPI",
           "GS10": "SH_GS10", "RATE GS10": "SH_GS10", "CAPE": "SH_CAPE",
           "TR CAPE": "SH_TRCAPE", "TR_CAPE": "SH_TRCAPE"}


def label(value) -> str:
    return " ".join(str(value if value is not None else "").upper().split())


def parse_rows(rows: list, now: str) -> dict:
    from pipeline.build import SOURCES, blank_series, normalize_observations
    header_index = next((i for i, row in enumerate(rows[:30])
                         if {"DATE", "P", "CPI"}.issubset({label(v) for v in row})), None)
    if header_index is None:
        raise ValueError("Shiller: unknown header layout")
    header = [label(v) for v in rows[header_index]]
    date_index = header.index("DATE")
    columns = {i: ALIASES[name] for i, name in enumerate(header) if name in ALIASES}
    if len(set(columns.values())) != len(columns):
        raise ValueError("Shiller: ambiguous duplicate columns")
    series = {s["id"]: blank_series(s) for s in SOURCES if s["id"] in columns.values()}
    notes = []
    for row in rows[header_index + 1:]:
        raw_date = row[date_index] if date_index < len(row) else None
        if raw_date in (None, ""):
            notes.extend(str(v).strip() for v in row if isinstance(v, str) and v.strip())
            continue
        if isinstance(raw_date, (date, datetime)):
            day = date(raw_date.year, raw_date.month, 1).isoformat()
        else:
            try:
                numeric = float(raw_date)
            except (ValueError, TypeError):
                continue
            if not math.isfinite(numeric):
                raise ValueError("Shiller: invalid date")
            year = int(numeric)
            month = round((numeric - year) * 100)
            if abs(numeric - (year + month / 100)) > 0.000001:
                raise ValueError("Shiller: expected YYYY.MM")
            day = date(year, month, 1).isoformat()
        if day > now[:10]:
            continue
        for index, key in columns.items():
            raw = row[index] if index < len(row) else None
            value = None if label(raw) in ("", "NA", "N/A", ".", "#N/A") else raw
            series[key]["observations"].append({"date": day, "value": value})
    for entry in series.values():
        entry["observations"] = normalize_observations(entry["observations"])
        if not any(p["value"] is not None for p in entry["observations"]):
            entry.update(reasonCode="download_failed", reason="El archivo Shiller no contiene valores para esta serie.")
            entry["observations"] = []
            continue
        entry.update(status="ok", fetchedAt=now)
        entry.pop("reason", None)
        entry.pop("reasonCode", None)
        if notes:
            entry["notes"] = "Notas del proveedor: " + " · ".join(dict.fromkeys(notes))
    return series


def parse_workbook(content: bytes, now: str) -> dict:
    if len(content) > MAX_BYTES:
        raise ValueError("Shiller: workbook too large")
    if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content)
        try:
            sheet = workbook.sheet_by_name("Data") if "Data" in workbook.sheet_names() else workbook.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            workbook.release_resources()
    elif content.startswith(b"PK\x03\x04"):
        import openpyxl
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * MAX_BYTES:
                raise ValueError("Shiller: expanded workbook too large")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.active
            rows = list(sheet.values)
        finally:
            workbook.close()
    else:
        raise ValueError("Shiller: response is not an Excel workbook")
    return parse_rows(rows, now)


def import_workbook(path: Path, now: str) -> dict:
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Shiller: workbook too large")
    return parse_workbook(path.read_bytes(), now)


class WorkbookLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.links.extend(value for key, value in attrs if key == "href" and value)


def workbook_url(html: str) -> str:
    parser = WorkbookLinks()
    parser.feed(html)
    candidates = set()
    for href in parser.links:
        url = urljoin(PAGE_URL, href)
        parts = urlsplit(url)
        if (parts.scheme == "https" and parts.hostname in {"shillerdata.com", "img1.wsimg.com"}
                and not parts.username and not parts.password and parts.port in (None, 443)
                and unquote(parts.path).split("/")[-1].lower() in {"ie_data.xls", "ie_data.xlsx"}):
            candidates.add(url)
    if len(candidates) != 1:
        raise ValueError("Shiller: official workbook link missing or ambiguous")
    return candidates.pop()


def fetch_bytes(client: httpx.Client, url: str) -> bytes:
    for attempt in range(3):
        try:
            with client.stream("GET", url) as response:
                response.raise_for_status()
                content = bytearray()
                for chunk in response.iter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_BYTES:
                        raise ValueError("Shiller: response too large")
                return bytes(content)
        except httpx.HTTPError as exc:
            transient = not isinstance(exc, httpx.HTTPStatusError) or exc.response.status_code == 429 or exc.response.status_code >= 500
            if not transient or attempt == 2:
                raise RuntimeError("Shiller: download failed") from None
            time.sleep(2 ** attempt)
    raise RuntimeError("Shiller: retries exhausted")


def download_workbook(now: str) -> dict:
    with httpx.Client(timeout=30, follow_redirects=False) as client:
        html = fetch_bytes(client, PAGE_URL).decode("utf-8")
        return parse_workbook(fetch_bytes(client, workbook_url(html)), now)


def acquire_shiller(now: str, previous: dict, *, approved: bool, local_file: str | None = None) -> dict:
    from pipeline.build import SOURCES, blank_series
    result = {s["id"]: blank_series(s) for s in SOURCES if s["provider"] == "Shiller"}
    if not approved:
        return result  # Withdrawal of permission also prevents reuse of previous data.
    try:
        imported = import_workbook(Path(local_file), now) if local_file else download_workbook(now)
        if set(imported) != set(result) or any(s["status"] != "ok" for s in imported.values()):
            raise ValueError("Shiller: incomplete workbook")
        return imported
    except Exception:
        # Excel readers raise several exception types. Isolate this provider, and
        # never print local paths, workbook contents or network request details.
        for key, entry in result.items():
            old = previous.get(key)
            if old and old["status"] in ("ok", "stale") and old["observations"]:
                entry.update(observations=copy.deepcopy(old["observations"]), fetchedAt=old["fetchedAt"],
                             status="stale", reason="Falló Shiller; se conserva la última descarga válida.")
                if old.get("notes"):
                    entry["notes"] = old["notes"]
            else:
                entry["reason"] = "No se pudo descargar o validar el archivo oficial Shiller."
            entry["reasonCode"] = "download_failed"
        print("WARNING Shiller: download or workbook validation failed; real fallback only")
        return result
