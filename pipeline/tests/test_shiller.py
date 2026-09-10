import copy
import io
from unittest.mock import Mock

import httpx
import openpyxl
import pytest

from pipeline import shiller
from pipeline.build import SOURCES, acquire_series, validate_snapshot

NOW = "2026-09-10T12:00:00Z"
HEADER = ["Date", "P", "D", "E", "CPI", "Fraction", "Rate GS10", "CAPE", "TR CAPE"]
ROWS = [["Multilevel heading"], HEADER,
        [1871.01, 4.44, .26, .4, 12.46, 1871.04, 5.32, "NA", "NA"],
        [2024.10, 5000, 70, 200, 310, 2024.83, 4, 30, 32],
        ["", "Latest price is provisional"]]


def test_original_headers_nulls_and_notes():
    parsed = shiller.parse_rows(ROWS, NOW)
    assert len(parsed) == 7
    assert parsed["SH_GS10"]["observations"][0]["value"] == 5.32
    assert parsed["SH_CAPE"]["observations"][0]["value"] is None
    assert parsed["SH_TRCAPE"]["observations"][-1]["date"] == "2024-10-01"
    assert all("provisional" in s["notes"] for s in parsed.values())
    assert all("reasonCode" not in s for s in parsed.values())
    validate_snapshot({"schemaVersion": 1, "mode": "live", "generatedAt": NOW,
                       "series": parsed, "events": [], "calendarStatus": "ok"})


def test_xlsx_without_positional_assumptions():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Data"
    for row in ROWS:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    assert shiller.parse_workbook(output.getvalue(), NOW)["SH_P"]["status"] == "ok"


@pytest.mark.parametrize("raw", [2024.13, 2024.001, float("inf")])
def test_invalid_dates_fail_closed(raw):
    with pytest.raises(ValueError):
        shiller.parse_rows([HEADER, [raw, 1, 1, 1, 1]], NOW)


def test_null_only_columns_are_not_successful():
    parsed = shiller.parse_rows([HEADER, [2024.10, 5, 0, 2, 300, 0, 4, "NA", "N/A"]], NOW)
    assert parsed["SH_D"]["observations"][0]["value"] == 0
    assert parsed["SH_CAPE"]["status"] == "unavailable"
    assert parsed["SH_CAPE"]["observations"] == []


def test_future_rows_are_excluded_and_duplicate_headers_rejected():
    result = shiller.parse_rows(ROWS + [[2027.01, 9, 1, 2, 300, 0, 4, 30, 30]], NOW)
    assert len(result["SH_P"]["observations"]) == 2
    with pytest.raises(ValueError):
        shiller.parse_rows([["Date", "P", "CPI", "P"]], NOW)


def test_permission_gate_prevents_reads_downloads_and_old_data(monkeypatch):
    download = Mock()
    read = Mock()
    monkeypatch.setattr(shiller, "download_workbook", download)
    monkeypatch.setattr(shiller, "import_workbook", read)
    previous = shiller.parse_rows(ROWS, NOW)
    result = shiller.acquire_shiller(NOW, previous, approved=False, local_file="ignored.xls")
    assert all(s["status"] == "unavailable" and not s["observations"] for s in result.values())
    assert all(s["reasonCode"] == "license_review" for s in result.values())
    download.assert_not_called()
    read.assert_not_called()


def test_authorized_workbook_success_and_fallback(monkeypatch):
    previous = shiller.parse_rows(ROWS, NOW)
    before = copy.deepcopy(previous)
    monkeypatch.setattr(shiller, "download_workbook", lambda _: previous)
    assert shiller.acquire_shiller(NOW, {}, approved=True) == previous
    monkeypatch.setattr(shiller, "download_workbook", Mock(side_effect=ValueError("bad workbook")))
    result = shiller.acquire_shiller("2026-09-11T12:00:00Z", previous, approved=True)
    assert all(s["status"] == "stale" and s["fetchedAt"] == NOW for s in result.values())
    assert all(s["reasonCode"] == "download_failed" for s in result.values())
    assert previous == before
    assert result["SH_P"]["observations"] is not previous["SH_P"]["observations"]


def test_partial_workbook_and_demo_fallback_rejected(monkeypatch):
    previous = shiller.parse_rows(ROWS, NOW)
    for series in previous.values():
        series["status"] = "demo"
    monkeypatch.setattr(shiller, "download_workbook", lambda _: {"SH_P": previous["SH_P"]})
    result = shiller.acquire_shiller(NOW, previous, approved=True)
    assert all(s["status"] == "unavailable" and not s["observations"] for s in result.values())


def test_official_link_is_discovered_and_other_hosts_rejected():
    url = "https://img1.wsimg.com/downloads/version/ie_data.xls?ver=123"
    assert shiller.workbook_url(f'<a href="{url}">Download</a>') == url
    for href in ["http://img1.wsimg.com/ie_data.xls", "https://evil.example/ie_data.xls",
                 "https://img1.wsimg.com@evil.example/ie_data.xls", "https://shillerdata.com:8080/ie_data.xls"]:
        with pytest.raises(ValueError):
            shiller.workbook_url(f'<a href="{href}">Download</a>')
    with pytest.raises(ValueError):
        shiller.parse_workbook(b"<html>not a workbook</html>", NOW)


def test_download_retries_and_size_limit(monkeypatch):
    monkeypatch.setattr(shiller.time, "sleep", lambda _: None)
    calls = []
    def respond(request):
        calls.append(request)
        return httpx.Response(503 if len(calls) < 3 else 200, content=b"ok")
    with httpx.Client(transport=httpx.MockTransport(respond)) as client:
        assert shiller.fetch_bytes(client, shiller.PAGE_URL) == b"ok"
    assert len(calls) == 3
    monkeypatch.setattr(shiller, "MAX_BYTES", 2)
    with httpx.Client(transport=httpx.MockTransport(lambda _: httpx.Response(200, content=b"large"))) as client:
        with pytest.raises(ValueError):
            shiller.fetch_bytes(client, shiller.PAGE_URL)


def test_fred_policy_and_attribution_survive_success(monkeypatch):
    from pipeline import build
    monkeypatch.setattr(build, "get_json", lambda *args: {"observations": [{"date": "2026-09-09", "value": "16"}]})
    vix = next(s for s in SOURCES if s["id"] == "VIXCLS")
    result = acquire_series(vix, Mock(), "test", NOW, {})
    assert vix["enabled"] and result["status"] == "ok"
    assert "Chicago Board Options Exchange" in result["attribution"]
    assert "reasonCode" not in result
    assert sum(s["enabled"] for s in SOURCES) == 64
    assert all(not s["enabled"] for s in SOURCES if s.get("reasonCode") == "permission_required")
